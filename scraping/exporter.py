"""Write results to a formatted Excel workbook."""
from datetime import datetime
from pathlib import Path

import pandas as pd
from django.conf import settings
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _style(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for idx, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 12), 60)


def write_excel(run_id, rows, tables):
    out_dir = Path(settings.MEDIA_ROOT)
    out_dir.mkdir(parents=True, exist_ok=True)
    name = f"scrape_{run_id}_{datetime.now():%Y-%m-%d_%H%M}.xlsx"

    df = pd.DataFrame(rows)
    if not df.empty:
        df = df.drop_duplicates()
    with pd.ExcelWriter(out_dir / name, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Results", index=False)
        _style(writer.sheets["Results"])
        for i, table in enumerate(tables[:50], start=1):
            sheet = f"Table {i}"
            table.to_excel(writer, sheet_name=sheet, index=False)
            _style(writer.sheets[sheet])
    return name, len(df)
